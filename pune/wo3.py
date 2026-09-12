import openpyxl,re,csv
FY={'2016-17':'xl/1617_1.xlsx','2017-18':'xl/1718_1.xlsx','2018-19':'xl/1819_2.xlsx','2019-20':'xl/1920_1.xlsx',
'2020-21':'xl/2021_1.xlsx','2021-22':'xl/2122_1.xlsx','2022-23':'xl/2223_1.xlsx','2023-24':'xl/2324_1.xlsx',
'2024-25':'xl/2425_1.xlsx','2025-26':'xl/2526_1.xlsx'}
DEPT=['roads_W2','electric_W1','buildings_W4','slum_W5','water_W6','sewerage_W7']
w=csv.writer(open('pmc_wardoffice.csv','w',newline=''))
w.writerow(['year','block','sr','prabhag_list','ward_office','*DEPT*']+DEPT+['total_lakh'])
for yr,f in FY.items():
    wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
    line=[yr]
    for pre in ['WardOfficePlan','WardOfficeNonPlan','CitizensPlan','CitizensNonPlan']:
        sh=[s for s in wb.sheetnames if s.startswith(pre)][0]
        rows=[[('' if c is None else str(c).strip()) for c in r] for r in wb[sh].iter_rows(values_only=True)]
        hi=next(i for i,r in enumerate(rows) if any(re.match(r'^[A-Z]{2}\d{2}[A-Z]\d{3}/[WC]\d+-$',c) for c in r))
        off=next(j for j,c in enumerate(rows[hi]) if re.match(r'^[A-Z]{2}\d{2}[A-Z]\d{3}/[WC]\d+-$',c))
        nm=off-1; wd=nm-1 if off>3 else 0
        n=0;tot=0
        for r in rows[hi+1:]:
            r=r+['']*(off+8)
            if not re.match(r'^\d+$',str(r[1])): 
                if n>0: break
                continue
            vals=[]
            for c in r[off:off+6]:
                try: vals.append(round(float(c),2))
                except: vals.append(0.0)
            t=round(sum(vals),2); tot+=t; n+=1
            w.writerow([yr,pre,r[1],r[wd],r[nm],'']+vals+[t])
        line.append(f"{pre[:14]}:{n}rows/{tot:,.0f}L")
    print(' | '.join(line)); wb.close()
